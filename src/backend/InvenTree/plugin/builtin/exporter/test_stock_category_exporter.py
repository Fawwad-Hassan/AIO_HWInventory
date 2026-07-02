"""Unit test for the stock category exporter plugin."""

import io

from django.urls import reverse

import openpyxl

from InvenTree.unit_test import InvenTreeAPITestCase
from plugin.registry import registry


class StockCategoryExporterTest(InvenTreeAPITestCase):
    """Test the multi-sheet (per-category) stock exporter plugin."""

    fixtures = ['category', 'part', 'location', 'stock', 'company']
    roles = ['stock.view']

    SLUG = 'inventree-stock-category-exporter'

    def setUp(self):
        """Create deterministic data: one serialized and one bulk category."""
        super().setUp()

        from part.models import Part, PartCategory
        from stock.models import StockItem

        # A category whose items are all serialized
        self.cat_serial = PartCategory.objects.create(name='ZZ Serialized Widgets')
        part_serial = Part.objects.create(
            name='ZZ Widget', category=self.cat_serial, component=True, trackable=True
        )
        StockItem.objects.create(part=part_serial, serial='SN-1', quantity=1)
        StockItem.objects.create(part=part_serial, serial='SN-2', quantity=1)

        # A category whose items are bulk (quantity only)
        self.cat_bulk = PartCategory.objects.create(name='ZZ Bulk Gadgets')
        part_bulk = Part.objects.create(
            name='ZZ Gadget', category=self.cat_bulk, component=True
        )
        StockItem.objects.create(part=part_bulk, quantity=25)

    def _header(self, worksheet) -> list:
        """Return the header row (first row) of a worksheet."""
        return [cell.value for cell in next(worksheet.iter_rows(max_row=1))]

    def test_multisheet_export(self):
        """Export stock as xlsx grouped by category and validate the workbook."""
        registry.set_plugin_state(self.SLUG, True)

        url = reverse('api-stock-list')

        with self.export_data(
            url, export_plugin=self.SLUG, export_format='xlsx', decode=False
        ) as data_file:
            self.assertIsInstance(data_file, io.BytesIO)

            workbook = openpyxl.load_workbook(data_file)

            # Our two deterministic categories each produce a sheet
            self.assertGreater(len(workbook.sheetnames), 1)
            self.assertIn('ZZ Serialized Widgets', workbook.sheetnames)
            self.assertIn('ZZ Bulk Gadgets', workbook.sheetnames)

            # Sheet names must be unique (case-insensitive) and <= 31 chars
            lowered = [name.lower() for name in workbook.sheetnames]
            self.assertEqual(len(lowered), len(set(lowered)))
            for name in workbook.sheetnames:
                self.assertLessEqual(len(name), 31)
                header = self._header(workbook[name])
                # Every sheet carries the curated columns
                self.assertIn('Asset Name', header)
                self.assertIn('Asset Type', header)
                self.assertIn('Location Name', header)
                # Curated set only: internal fields must not leak through
                self.assertNotIn('Barcode Hash', header)

            # Serialized category: Serial Number column, no Quantity column
            serial_header = self._header(workbook['ZZ Serialized Widgets'])
            self.assertIn('Serial Number', serial_header)
            self.assertNotIn('Quantity', serial_header)

            # Bulk category: Quantity column, no Serial Number column
            bulk_header = self._header(workbook['ZZ Bulk Gadgets'])
            self.assertIn('Quantity', bulk_header)
            self.assertNotIn('Serial Number', bulk_header)

        # Selecting the plugin with a non-xlsx format must fall back cleanly
        # to a normal single-sheet file (still including the Category column).
        with self.export_data(
            url, export_plugin=self.SLUG, export_format='csv'
        ) as data_file:
            self.process_csv(data_file, required_cols=['Asset Name', 'Asset Type'])

        registry.set_plugin_state(self.SLUG, False)

    def test_multi_value_export_plugin_param(self):
        """Multi-sheet triggers even when export_plugin is sent more than once.

        The frontend sends export_plugin twice (a default, then the user's
        selection), so the value reaches the serializer as a list. The last
        value must win, matching how the view resolves the plugin.
        """
        from collections import OrderedDict

        from stock.exporter_helpers import CATEGORY_ID_KEY, CATEGORY_NAME_KEY
        from stock.serializers import StockItemSerializer

        serializer = StockItemSerializer(
            exporting=True,
            context={
                'export_plugin': ['inventree-exporter', self.SLUG]  # default, then ours
            },
        )
        headers = OrderedDict([
            ('part_detail.name', 'Asset Name'),
            (CATEGORY_NAME_KEY, 'Asset Type'),
        ])
        data = [
            {
                'part_detail': {'name': 'A'},
                CATEGORY_ID_KEY: 1,
                CATEGORY_NAME_KEY: 'Cat A',
            },
            {
                'part_detail': {'name': 'B'},
                CATEGORY_ID_KEY: 2,
                CATEGORY_NAME_KEY: 'Cat B',
            },
        ]

        result = serializer.export_to_file(data, headers, 'xlsx')
        workbook = openpyxl.load_workbook(io.BytesIO(result))
        self.assertEqual(set(workbook.sheetnames), {'Cat A', 'Cat B'})
